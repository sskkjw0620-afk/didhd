import sys
import os
import json
from pptx import Presentation


def save_image(image, out_path):
    with open(out_path, 'wb') as f:
        f.write(image.blob)


def extract(pptx_path, out_dir):
    prs = Presentation(pptx_path)
    slides_out = []
    assets_dir = os.path.join(out_dir, 'public', 'assets')
    os.makedirs(assets_dir, exist_ok=True)

    for i, slide in enumerate(prs.slides, start=1):
        slide_obj = {'index': i, 'texts': [], 'images': []}

        # extract text from shapes
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                text = shape.text.strip()
                if text:
                    slide_obj['texts'].append(text)

        # extract images from shapes
        img_count = 0
        for shape in slide.shapes:
            if shape.shape_type == 13 and hasattr(shape, 'image'):
                img = shape.image
                img_ext = img.ext
                img_count += 1
                img_name = f"slide{i}_img{img_count}.{img_ext}"
                img_path = os.path.join(assets_dir, img_name)
                save_image(img, img_path)
                slide_obj['images'].append(f"/assets/{img_name}")

        slides_out.append(slide_obj)

    # write slides.json into public
    public_dir = os.path.join(out_dir, 'public')
    os.makedirs(public_dir, exist_ok=True)
    with open(os.path.join(public_dir, 'slides.json'), 'w', encoding='utf-8') as f:
        json.dump(slides_out, f, ensure_ascii=False, indent=2)

    print(f"Extracted {len(slides_out)} slides to {public_dir}")


def main():
    if len(sys.argv) < 3:
        print("Usage: python extract_pptx.py path/to/input.pptx path/to/output_folder")
        sys.exit(1)

    pptx_path = sys.argv[1]
    out_dir = sys.argv[2]
    if not os.path.exists(pptx_path):
        print("PPTX file not found:", pptx_path)
        sys.exit(1)

    extract(pptx_path, out_dir)


if __name__ == '__main__':
    main()
